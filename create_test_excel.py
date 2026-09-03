"""
Script para crear un archivo Excel de prueba con el formato esperado.
"""
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Horarios"

# Encabezados
headers = ["Piso", "Ambiente", "Sesión", "Hora Inicio", "Hora Fin",
           "Docente Líder", "Docentes Colaboradores", "Pacientes Estandarizados"]

for col, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=header)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="2D3154", end_color="2D3154", fill_type="solid")
    cell.alignment = Alignment(horizontal="center")

# Datos de prueba
data = [
    # Piso, Ambiente, Sesión, Inicio, Fin, Docente, Colaboradores, Pacientes
    # ── MEDICINA HUMANA (NARANJA) ──
    ["Primer Piso", "SL02LA25: SALA DE HABILIDADES 4",
     "[PAX] - VALORACION DE LA ESCENA Y CINEMATICA DEL TRAUMA",
     "07:00", "08:40", "Leslie Isabel Barzola Bueno", "Ninguno", "Ninguno"],
    ["Primer Piso", "SALA ALTA FIDEL OPERACIONES",
     "[SimQx] - TALLER DE CIRUGÍA GENERAL",
     "09:00", "11:30", "Dr. Carlos Mendoza Ríos", "Dra. Ana García", "Ninguno"],
    ["Primer Piso", "SALA ALTA FIDEL SHOCK TRAUMA",
     "[SimPed] - MANEJO DE URGENCIAS PEDIÁTRICAS",
     "07:30", "09:30", "Dr. Fernando Castro", "Ninguno", "1 maniquí"],
    ["Primer Piso", "SALA ALTA FIDEL 4-Parto",
     "[SimGO] - MANEJO DE EMERGENCIAS GINECO-OBSTÉTRICAS",
     "10:00", "12:30", "Dra. Carla Ramos", "Dr. José Torres", "2 pacientes"],
    ["Primer Piso", "SALA DE HOSPITALIZACIÓN 1",
     "[SCI] - SOPORTE CLÍNICO INICIAL",
     "13:00", "15:00", "Dr. Roberto Silva", "Dra. Paula Vega", "Ninguno"],
    ["Primer Piso", "LAB. TECNOLOGÍA MÉDICA",
     "[MED SBS] - SISTEMAS BIOLÓGICOS Y SALUD",
     "08:00", "10:00", "Dra. María López", "Ninguno", "Ninguno"],
    ["Primer Piso", "CONSULTORIO 01",
     "[Ext Med] - MEDICINA INTERNA HOSPITALARIA",
     "14:00", "16:00", "Dr. Hugo Ramírez", "Ninguno", "Ninguno"],
    ["Primer Piso", "CONSULTORIO 02",
     "[Ext Cirg] - TALLER PRÁCTICO DE CIRUGÍA",
     "16:30", "18:30", "Dr. Mario Benítez", "Ninguno", "Ninguno"],

    # ── ENFERMERÍA (CELESTE) ──
    ["Segundo Piso", "SALA DE HABILIDADES 1",
     "[ENF] - CUIDADOS DE ENFERMERÍA EN EL ADULTO",
     "07:00", "09:00", "Lic. Lucía Flores", "Lic. Pedro Guzmán", "Ninguno"],
    ["Segundo Piso", "SALA DE HABILIDADES 2",
     "[ENF] - ADMINISTRACIÓN DE MEDICAMENTOS Y VENOCLISIS",
     "09:30", "12:00", "Lic. Ana Martínez", "Lic. Carlos Pérez", "Ninguno"],

    # ── OBSTETRICIA (VIOLETA) ──
    ["Segundo Piso", "SALA DE HABILIDADES 3",
     "[OBST] - ATENCIÓN DEL PARTO Y ALUMBRAMIENTO",
     "07:00", "09:30", "Lic. Carmen Vidal", "Lic. Ricardo Salas", "3 pacientes"],
    ["Segundo Piso", "CONSULTORIO 05",
     "[OBST] - CONTROL PRENATAL Y PSICOPROFILAXIS",
     "10:00", "12:00", "Lic. Patricia Nava", "Ninguno", "Ninguno"],

    # ── MÁS CURSOS DE MEDICINA EN PISO 2 ──
    ["Segundo Piso", "SALA DE ECOGRAFÍAS",
     "ECOGRAFIA - EVALUACIÓN ULTRASONOGRÁFICA FAST",
     "07:00", "09:00", "Dra. Sofía Ruiz", "Dr. Miguel León", "Ninguno"],
    ["Segundo Piso", "CONSULTORIO 06",
     "[Ext Ped] - PEDIATRÍA CLÍNICA",
     "13:00", "15:00", "Dr. Javier Mora", "Ninguno", "Ninguno"],
    ["Segundo Piso", "CONSULTORIO 07",
     "[Ext GyO] - GINECOLOGÍA Y OBSTETRICIA",
     "15:30", "17:30", "Dra. Claudia Viteri", "Ninguno", "Ninguno"],
]


for row_idx, row_data in enumerate(data, 2):
    for col_idx, value in enumerate(row_data, 1):
        ws.cell(row=row_idx, column=col_idx, value=value)

# Ajustar anchos de columna
ws.column_dimensions['A'].width = 15
ws.column_dimensions['B'].width = 40
ws.column_dimensions['C'].width = 55
ws.column_dimensions['D'].width = 12
ws.column_dimensions['E'].width = 12
ws.column_dimensions['F'].width = 25
ws.column_dimensions['G'].width = 25
ws.column_dimensions['H'].width = 20

output_path = os.path.join(os.path.dirname(__file__), "datos_prueba.xlsx")
wb.save(output_path)
print(f"Excel de prueba creado: {output_path}")
print(f"Total filas de datos: {len(data)}")
